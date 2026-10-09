#!/usr/bin/env python3
"""Reads a .docx with python-docx and prints structural facts as JSON: the independent reading of a document that
tests compare converter output against. Usage: docx_facts.py FILE"""
import json
import sys

import docx
from docx.oxml.ns import qn


def main(path):
    document = docx.Document(path)
    headings = []
    for paragraph in document.paragraphs:
        name = paragraph.style.name
        if name == 'Title' or name.startswith('Heading'):
            headings.append({'style': name, 'text': paragraph.text})
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                tc = cell._tc
                props = tc.tcPr
                span = props.find(qn('w:gridSpan')) if props is not None else None
                cells.append(cell.text + ('|span=%s' % span.get(qn('w:val')) if span is not None else ''))
            rows.append(cells)
        tables.append(rows)
    print(
        json.dumps(
            {
                'paragraphs': len(document.paragraphs),
                'headings': headings,
                'tables': tables,
                'inlineShapes': len(document.inline_shapes),
                'text': [p.text for p in document.paragraphs],
            }
        )
    )


if __name__ == '__main__':
    main(sys.argv[1])
